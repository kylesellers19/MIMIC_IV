#!/usr/bin/env python3
"""Adjudication instrument — a blank .xlsx form with dropdown validation, one row per blinded ID.

Sheets: 'sample' (R0001–R0400 in master-key order, never reordered by cell), 'training' (T001–T020),
'lists' (dropdown options), 'README' (pointer to the data dictionary). Columns exactly as criteria §7:
review_id; primary_response; temporality; q6a_cardiometabolic; q6b_alcohol; q6c_other_cause; q6c_specify;
rationale; negation_override; minutes_spent; reviewer_initials.

The workbook contains ONLY blinded IDs — no cell, no weight, no double_review flag, no hadm_id — so the
tracked template at docs/adjudication/adjudication_template.xlsx is commit-safe. `build_form()` is reused by
scripts/assign_reviewers.py to write the per-reviewer forms (untracked, under data/processed/adjudication/).

Usage: scripts/run0919_env.sh scripts/make_adjudication_template.py [--out docs/adjudication/adjudication_template.xlsx]
"""
import argparse
import os
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.getcwd())
from scripts.masld_config import TRAINING_KEY

MASTER_KEY = "data/processed/aim1a_sample.parquet"
DEFAULT_OUT = "docs/adjudication/adjudication_template.xlsx"

COLUMNS = ["review_id", "primary_response", "temporality", "q6a_cardiometabolic", "q6b_alcohol",
           "q6c_other_cause", "q6c_specify", "rationale", "negation_override", "minutes_spent", "reviewer_initials"]

OPTIONS = {
    "primary_response": ["Definite", "Probable", "Possible", "Absent", "Insufficient information"],
    "temporality": ["During this admission", "Before this admission", "Both", "Not applicable"],
    "q6a_cardiometabolic": ["Yes", "No", "Insufficient", "Not applicable"],
    "q6b_alcohol": ["(i) Below MetALD threshold", "(ii) Within MetALD band (140-350 g/wk F, 210-420 g/wk M)",
                    "(iii) Above MetALD band", "(iv) Not documented / insufficient", "Not applicable"],
    "q6c_other_cause": ["Yes (specify)", "No", "Unclear", "Not applicable"],
    "negation_override": ["Yes", "No"],
}
WIDTHS = {"review_id": 11, "primary_response": 24, "temporality": 24, "q6a_cardiometabolic": 20, "q6b_alcohol": 44,
          "q6c_other_cause": 17, "q6c_specify": 28, "rationale": 60, "negation_override": 18, "minutes_spent": 14, "reviewer_initials": 17}
HEADER_FILL = PatternFill("solid", fgColor="DDDDDD")


def _lists_sheet(wb):
    ws = wb.create_sheet("lists")
    ranges = {}
    for j, (col, opts) in enumerate(OPTIONS.items(), start=1):
        ws.cell(row=1, column=j, value=col).font = Font(bold=True)
        for i, o in enumerate(opts, start=2):
            ws.cell(row=i, column=j, value=o)
        L = get_column_letter(j)
        ranges[col] = f"=lists!${L}$2:${L}${len(opts) + 1}"
        ws.column_dimensions[L].width = max(18, max(len(o) for o in opts) + 2)
    return ranges


def _form_sheet(wb, name, ids, ranges, id_header="review_id"):
    ws = wb.create_sheet(name)
    for j, col in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=j, value=(id_header if col == "review_id" else col))
        c.font = Font(bold=True)
        c.fill = HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(j)].width = WIDTHS[col]
    n = len(ids)
    for i, rid in enumerate(ids, start=2):
        c = ws.cell(row=i, column=1, value=rid)
        c.protection = Protection(locked=True)
        for j in range(2, len(COLUMNS) + 1):
            cell = ws.cell(row=i, column=j)
            cell.protection = Protection(locked=False)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    if n:
        for col, ref in ranges.items():
            j = COLUMNS.index(col) + 1
            L = get_column_letter(j)
            dv = DataValidation(type="list", formula1=ref, allow_blank=True, showErrorMessage=True,
                                errorTitle="Choose from the list", error="Please pick one of the listed responses.")
            ws.add_data_validation(dv)
            dv.add(f"{L}2:{L}{n + 1}")
        j = COLUMNS.index("minutes_spent") + 1
        L = get_column_letter(j)
        dv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True, showErrorMessage=True,
                            errorTitle="Minutes", error="Whole number of minutes (0 or more).")
        ws.add_data_validation(dv)
        dv.add(f"{L}2:{L}{n + 1}")
    ws.freeze_panes = "B2"
    # protect the ID column only; everything else stays editable, and the sheet can be unprotected without a password
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.sort = False
    ws.protection.autoFilter = False
    ws.protection.formatCells = False
    return ws


def _readme(wb):
    ws = wb.create_sheet("README", 0)
    lines = [
        "Hepatic steatosis chart adjudication form",
        "",
        "One row per chart, identified only by its blinded review ID. Fill the dropdowns; free text in q6c_specify and rationale.",
        "Column definitions and the response scale: docs/adjudication/data_dictionary.md; criteria: docs/adjudication_criteria.md.",
        "Sheet 'sample' = formal review; sheet 'training' = the 20-chart calibration round (T001–T020).",
        "rationale is REQUIRED when primary_response is 'Possible' and whenever negation_override is 'Yes'.",
        "Questions 6a–6c are asked only when primary_response is Definite or Probable; otherwise choose 'Not applicable'.",
        "minutes_spent: whole minutes on the chart, including reading and form entry.",
        "Do not add, delete or reorder rows. The ID column is locked (the sheet protection has no password).",
        "Do not paste any chart text into this form.",
    ]
    for i, t in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=t)
        if i == 1:
            c.font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 120
    return ws


def build_form(sample_ids, training_ids, out_path):
    """Write a blank form for the given (already ordered) blinded IDs. Returns the path."""
    wb = Workbook()
    wb.remove(wb.active)
    _readme(wb)
    ranges = _lists_sheet(wb)
    _form_sheet(wb, "sample", list(sample_ids), ranges)
    _form_sheet(wb, "training", list(training_ids), ranges)
    wb.move_sheet("lists", offset=len(wb.sheetnames))   # lists last
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)
    return out_path


def load_ids(master_key=MASTER_KEY, training_key=TRAINING_KEY):
    sample_ids = pd.read_parquet(master_key, columns=["review_id"])["review_id"].tolist()   # master-key order
    training_ids = pd.read_parquet(training_key, columns=["training_id"])["training_id"].tolist()
    assert len(sample_ids) == 400 and len(set(sample_ids)) == 400 and all(r.startswith("R") for r in sample_ids)
    assert len(training_ids) == 20 and all(t.startswith("T") for t in training_ids)
    return sample_ids, training_ids


def verify_commit_safe(path):
    """The workbook must contain only the ID columns' blinded IDs and static text: no digits-only cells > 4 digits."""
    from openpyxl import load_workbook
    wb = load_workbook(path)
    n_cells = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if v is None:
                    continue
                n_cells += 1
                s = str(v)
                assert not (s.isdigit() and len(s) > 4), "unexpected long integer in template"
                assert "code+/note" not in s and "code-/note" not in s and "double" not in s.lower(), "cell/double-review leakage"
    return n_cells


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=DEFAULT_OUT)
    a = ap.parse_args()
    sample_ids, training_ids = load_ids()
    build_form(sample_ids, training_ids, a.out)
    n = verify_commit_safe(a.out)
    print(f"wrote {a.out}: sample rows {len(sample_ids)}, training rows {len(training_ids)}, non-empty cells {n}, size {os.path.getsize(a.out):,} bytes")


if __name__ == "__main__":
    main()
